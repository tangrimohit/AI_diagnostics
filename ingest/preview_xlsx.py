"""One-off CLI to peek at an xlsx's structure before ingesting it."""
from __future__ import annotations
import sys
from pathlib import Path
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def preview(path: Path, sample_rows: int = 3) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    print(f"\n=== {path.name} ===")
    print(f"Sheets: {wb.sheetnames}\n")
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        print(f"--- Sheet: {name} ---")
        print(f"  dims: {ws.max_row} rows x {ws.max_column} cols")
        if not rows:
            print("  (empty)\n")
            continue
        header = rows[0]
        print(f"  headers: {list(header)}")
        for i, row in enumerate(rows[1 : 1 + sample_rows]):
            print(f"  row {i + 1}: {list(row)}")
        print()
    wb.close()


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        preview(Path(arg))
